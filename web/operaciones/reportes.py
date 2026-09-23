import io
import unicodedata

import streamlit as st
from datetime import datetime
from core.database import get_db_session
from core.models import Cuenta, Celula, Cliente, Tarea, AlertaHistorial
from core.registro import obtener_acciones
from web.ui import cabecera, stat, divider


# Columnas de la "tabla de exitos" (mismo orden en pantalla y en el Excel).
COLUMNAS_EXITOS = ("Fecha", "Usuario", "Tipo", "Resultado", "Link")

# Resultado por rol canonico: (etiqueta con link, etiqueta base sin link).
_ETIQUETAS_RESULTADO = {
    "cita": ("Cita (RT con cita)", "Cita"),
    "hashtags": ("Post", "Post"),
    "comentario": ("Comentario", "Comentario"),
}

# Roles que NO generan link publicable: etiqueta fija y link SIEMPRE vacio.
_RESULTADOS_FIJOS = {
    "rt": "RT simple (no genera link)",
    "like": "Like (no genera link)",
}

_SUFIJO_SIN_LINK = " (link no capturado)"


def _es_url_post(url) -> bool:
    """True solo si `url` es un enlace http(s) a un POST (contiene /status/).

    Descarta perfiles ("https://twitter.com/usuario"), cadenas vacias, None,
    anchors y basura sin esquema http(s). Nunca lanza.
    """
    try:
        texto = str(url or "").strip()
        if not texto:
            return False
        bajo = texto.lower()
        if not bajo.startswith(("http://", "https://")):
            return False
        if "/status/" not in bajo:
            return False
        # Debe quedar un id tras "/status/" (descarta ".../status/" pelado).
        resto = bajo.split("/status/", 1)[1]
        return bool(resto.strip(" /?#\t\r\n"))
    except Exception:
        return False


def _normalizar_tipo_local(tipo) -> str:
    """Fallback local de `core.registro.normalizar_rol_cuota`.

    Minusculas, sin acentos ni guiones; devuelve el rol canonico
    (hashtags/cita/rt/comentario/like) o el texto limpio si no lo reconoce.
    """
    texto = str(tipo or "").strip().lower()
    for separador in ("_", "-", ".", "/", ",", ";", "(", ")", "[", "]"):
        texto = texto.replace(separador, " ")
    texto = "".join(
        caracter
        for caracter in unicodedata.normalize("NFKD", texto)
        if not unicodedata.combining(caracter)
    )
    tokens = texto.split()
    if not tokens:
        return ""
    if any(token.startswith("cita") or token.startswith("quote") for token in tokens):
        return "cita"
    if any(token.startswith("hashtag") or token.startswith("mencion") for token in tokens):
        return "hashtags"
    if any(
        token
        in (
            "post",
            "posts",
            "publicacion",
            "publicaciones",
            "mantenimiento",
            "calentamiento",
            "hilo",
            "hilos",
        )
        or token.startswith("publicac")
        for token in tokens
    ):
        return "hashtags"
    if any(
        token.startswith("coment")
        or token in ("reply", "responder", "respuesta", "respuestas")
        for token in tokens
    ):
        return "comentario"
    if any(
        token == "rt" or token.startswith("retweet") or token.startswith("repost")
        for token in tokens
    ):
        return "rt"
    if any(token in ("like", "likes", "megusta") for token in tokens):
        return "like"
    return " ".join(tokens)


def _normalizar_tipo(tipo) -> str:
    """Rol canonico de un tipo de accion (`core.registro` o fallback local)."""
    try:
        from core.registro import normalizar_rol_cuota

        return normalizar_rol_cuota(tipo)
    except Exception:
        return _normalizar_tipo_local(tipo)


def _resultado_accion(tipo, url) -> tuple:
    """Devuelve `(resultado, link)` para una accion de la tabla de exitos.

    - cita/quote -> "Cita (RT con cita)"; post/hashtags/publicacion/
      mantenimiento/calentamiento/hilo -> "Post"; comentario/respuesta/reply ->
      "Comentario". Si la categoria genera link y `url` es un post valido
      (`_es_url_post`) se devuelve la URL; si no, "(link no capturado)".
    - rt/retweet/repost -> "RT simple (no genera link)" y link SIEMPRE "".
    - like -> "Like (no genera link)" y link SIEMPRE "" (aunque la URL
      historica apunte al tweet likeado).
    - Tipo desconocido -> el tipo tal cual (o "Accion") y, sin link,
      "(link no capturado)".

    Nunca lanza.
    """
    try:
        link = str(url or "").strip() if _es_url_post(url) else ""
        rol = _normalizar_tipo(tipo)
        if rol in _RESULTADOS_FIJOS:
            return _RESULTADOS_FIJOS[rol], ""
        if rol in _ETIQUETAS_RESULTADO:
            etiqueta, base = _ETIQUETAS_RESULTADO[rol]
            return (etiqueta, link) if link else (base + _SUFIJO_SIN_LINK, "")
        base = str(tipo or "").strip() or "Acción"
        return (base, link) if link else (base + _SUFIJO_SIN_LINK, "")
    except Exception:
        return (str(tipo or "").strip() or "Acción", "")


def _es_fila_vacia(fila: dict) -> bool:
    """True si TODOS los valores de la fila estan vacios (fila separadora)."""
    if not isinstance(fila, dict):
        return True
    return not any(str(valor or "").strip() for valor in fila.values())


def _filas_con_separacion(filas: list[dict]) -> list[dict]:
    """Intercala UNA fila vacia entre cada exito y el siguiente.

    Orden resultante: exito1, vacia, exito2, vacia, ..., exitoN. Sin fila
    vacia antes del primero ni despues del ultimo; con 0 o 1 exitos no hay
    ninguna vacia. Las filas separadoras tienen TODAS sus claves en "" (las
    mismas claves que usan los exitos). Es idempotente: si la lista ya viene
    separada, vuelve a dejar una sola fila vacia entre exitos.

    Args:
        filas: lista de dicts con los exitos ("Fecha", "Usuario", "Tipo",
            "Resultado", "Link"). Nunca lanza.

    Returns:
        Nueva lista (no modifica la original) con los exitos en orden y las
        filas vacias intercaladas.
    """
    limpias = [
        fila
        for fila in (filas or [])
        if isinstance(fila, dict) and not _es_fila_vacia(fila)
    ]
    if not limpias:
        return []

    claves: list = []
    for fila in limpias:
        for clave in fila:
            if clave not in claves:
                claves.append(clave)

    vacia = {clave: "" for clave in claves}
    separadas: list = []
    for indice, fila in enumerate(limpias):
        if indice:
            separadas.append(dict(vacia))
        separadas.append(fila)
    return separadas


def _excel_exitos_bytes(filas: list[dict]) -> bytes:
    """Genera EN MEMORIA el .xlsx de exitos con una fila vacia en medio.

    Fila 1: cabeceras con estilo (negrita, fondo 1DA1F2 y letra blanca) y
    ancho 22 por columna. Luego las filas tal cual: un exito, una fila
    completamente vacia, el siguiente exito, etc. Las filas vacias se escriben
    sin valores (todas las celdas vacias).

    Args:
        filas: exitos crudos o ya separados (idempotente); se aplica
            `_filas_con_separacion` antes de escribir.

    Returns:
        Los bytes del .xlsx, o b"" si openpyxl falla (nunca lanza).
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        separadas = _filas_con_separacion(filas)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Exitos"
        ws.append(list(COLUMNAS_EXITOS))

        for columna in range(1, len(COLUMNAS_EXITOS) + 1):
            celda = ws.cell(row=1, column=columna)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1DA1F2")
            celda.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(columna)].width = 22

        for fila in separadas:
            ws.append([
                "" if fila.get(columna) is None else fila.get(columna, "")
                for columna in COLUMNAS_EXITOS
            ])

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    except Exception:
        return b""


def render(usuario: dict):
    cabecera("📊 REPORTES", "Estadísticas generales del sistema")
    
    st.markdown("## Resumen General")
    
    with get_db_session() as db:
        total_cuentas = db.query(Cuenta).count()
        activas = db.query(Cuenta).filter(Cuenta.activa == True).count()
        celulas = db.query(Celula).filter(Celula.activa == True).count()
        clientes = db.query(Cliente).filter(Cliente.activo == True).count()
        tareas_total = db.query(Tarea).count()
        tareas_pend = db.query(Tarea).filter(Tarea.estado == "pendiente").count()
        alertas = db.query(AlertaHistorial).count()
        alertas_hoy = db.query(AlertaHistorial).filter(
            AlertaHistorial.fecha_envio >= datetime.now().replace(hour=0, minute=0)
        ).count()
        posts_hoy = db.query(Tarea).filter(
            Tarea.tipo == "post",
            Tarea.estado == "completada",
            Tarea.fecha_hora >= datetime.now().replace(hour=0, minute=0)
        ).count()
    
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        stat(f"{activas}/{total_cuentas}", "Cuentas activas")
    with c2:
        stat(celulas, "Células")
    with c3:
        stat(clientes, "Clientes")
    with c4:
        stat(tareas_pend, "Tareas pendientes")
    
    divider()
    
    c1, c2, c3 = st.columns(3)
    with c1:
        stat(posts_hoy, "Posts hoy")
    with c2:
        stat(alertas_hoy, "Alertas hoy")
    with c3:
        stat(alertas, "Alertas totales")
    
    divider()
    
    st.markdown("## Cuentas por Plataforma")
    
    with get_db_session() as db:
        por_plataforma = {}
        for c in db.query(Cuenta).all():
            por_plataforma.setdefault(c.plataforma, {"total": 0, "activas": 0})
            por_plataforma[c.plataforma]["total"] += 1
            if c.activa:
                por_plataforma[c.plataforma]["activas"] += 1
    
    if por_plataforma:
        for plat, datos in por_plataforma.items():
            st.markdown(f"**{plat}**: {datos['activas']}/{datos['total']} activas")
            st.progress(datos["activas"] / datos["total"] if datos["total"] else 0)

    divider()

    st.markdown("## Historial de acciones (mantenimientos / activaciones)")
    try:
        acciones = obtener_acciones(limit=50, solo_exitosas=True)
    except TypeError:
        # Compatibilidad mientras core/registro.py no tenga el parámetro:
        estados_ok = {"exito", "exitoso", "ok"}
        try:
            acciones = [
                a for a in obtener_acciones(limit=200)
                if (getattr(a, "estado", "") or "").strip().lower() in estados_ok
            ][:50]
        except Exception:
            acciones = []
    except Exception:
        acciones = []
    if acciones:
        filas = []
        for a in acciones:
            tipo = getattr(a, "tipo", "") or ""
            resultado, link = _resultado_accion(
                tipo, getattr(a, "url_publicacion", "") or ""
            )
            fecha = getattr(a, "fecha", None)
            filas.append({
                "Fecha": fecha.strftime("%Y-%m-%d %H:%M") if fecha else "",
                "Usuario": getattr(a, "usuario", "") or "",
                "Tipo": tipo,
                "Resultado": resultado,
                "Link": link,
            })
        filas_mostradas = _filas_con_separacion(filas)
        st.caption("✅ Solo se muestran las acciones exitosas; las fallidas no se incluyen.")
        st.dataframe(filas_mostradas, use_container_width=True, column_config={
            "Link": st.column_config.LinkColumn("Link", display_text="🔗 Abrir"),
        })
        datos_excel = _excel_exitos_bytes(filas_mostradas)
        if datos_excel:
            st.download_button(
                "⬇️ Descargar Excel (.xlsx)",
                data=datos_excel,
                file_name=f"exitos_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="dl_reportes_exitos",
            )
        else:
            st.warning(
                "No se pudo generar el Excel (openpyxl no disponible). "
                "Puedes descargar el CSV desde el menú del dataframe."
            )
    else:
        st.info("Aún no hay acciones registradas.")